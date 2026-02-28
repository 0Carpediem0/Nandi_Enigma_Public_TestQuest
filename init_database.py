import argparse
import csv
import getpass
import os
import sys
import traceback
from pathlib import Path
from typing import Iterable, Optional

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

# Подгрузка backend/.env, чтобы PGHOST, PGPORT, PGUSER, PGPASSWORD, PGDATABASE были заданы
_env_path = Path(__file__).resolve().parent / "backend" / ".env"
try:
    from dotenv import load_dotenv
    if _env_path.exists():
        load_dotenv(_env_path, encoding="utf-8")
    load_dotenv(encoding="utf-8")
except ImportError:
    pass

import psycopg
from psycopg import Connection as PgConnection
from psycopg import sql


_project_root = Path(__file__).resolve().parent
KB_XLSX_PATH = _project_root / "kb_test.xlsx"
if not KB_XLSX_PATH.exists():
    KB_XLSX_PATH = _project_root.parent / "kb_test.xlsx"


def _normalize_header(h: Optional[str]) -> str:
    return (h or "").strip().lower().replace(" ", "_")


def _seed_kb_from_xlsx(conn: PgConnection) -> None:
    """Заполняет knowledge_base из kb_test.xlsx: очищает старую вставку и вставляет строки из файла.
    В колонку tags попадают шаблонные вопросы — примеры запросов пользователя, которые могут
    привести к этой теме (например: вопрос «что даёт корова?» → тема «молоко»).
    """
    print("\n== База знаний из kb_test.xlsx ==")
    print(f"Путь к файлу: {KB_XLSX_PATH}")
    if not KB_XLSX_PATH.exists():
        print(f"Файл не найден: {KB_XLSX_PATH}")
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM knowledge_base")
            print(f"Записей в knowledge_base: {cur.fetchone()[0]}")
        return
    if not openpyxl:
        print("Установите openpyxl для чтения xlsx: pip install openpyxl")
        return

    wb = openpyxl.load_workbook(KB_XLSX_PATH, read_only=True, data_only=True)
    ws = wb.active
    if not ws:
        wb.close()
        print("В книге нет активного листа.")
        return
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        wb.close()
        print("Файл пустой.")
        return
    headers = [_normalize_header(str(h)) for h in header_row]
    # Маппинг возможных названий колонок -> поля БД
    # tags в БД = шаблонные вопросы (примеры запросов пользователя, которые приводят к этой теме)
    title_keys = ("question_template", "question", "вопрос", "title", "заголовок")
    content_keys = ("answer_template", "answer", "ответ", "content", "содержание", "текст")
    cat_key = "category"
    if "категория" in headers:
        cat_key = "категория"
    # Колонка с шаблонными вопросами (приводящими к этой теме); при отсутствии используем title
    template_questions_keys = ("шаблонные_вопросы", "template_questions", "примеры_вопросов", "теги", "tags")

    def col_index(keys: tuple) -> int:
        for k in keys:
            if k in headers:
                return headers.index(k)
        return -1

    title_idx = col_index(title_keys)
    content_idx = col_index(content_keys)
    cat_idx = headers.index(cat_key) if cat_key in headers else -1
    template_questions_idx = col_index(template_questions_keys)

    if title_idx < 0 or content_idx < 0:
        wb.close()
        print("В xlsx нужны колонки для вопроса (question_template/question/вопрос/title) и ответа (answer_template/answer/ответ/content).")
        return

    rows = []
    for row in rows_iter:
        if not row or len(row) <= max(title_idx, content_idx):
            continue
        title_val = row[title_idx]
        content_val = row[content_idx]
        title = (title_val if title_val is not None else "").strip() if isinstance(title_val, str) else str(title_val or "").strip()
        content = (content_val if content_val is not None else "").strip() if isinstance(content_val, str) else str(content_val or "").strip()
        if not title or not content:
            continue
        short = content[:500] if len(content) > 500 else content
        cat = None
        if cat_idx >= 0 and cat_idx < len(row) and row[cat_idx] is not None:
            cat = str(row[cat_idx]).strip() or None
        # Шаблонные вопросы: фразы, которые пользователь мог бы спросить и попасть на эту тему
        tags_list = []
        if template_questions_idx >= 0 and template_questions_idx < len(row) and row[template_questions_idx]:
            raw = str(row[template_questions_idx]).strip()
            for s in raw.replace("|", ";").replace(",", ";").split(";"):
                t = s.strip()
                if t:
                    tags_list.append(t)
        if not tags_list:
            tags_list = [title]
        rows.append((title, content, short, tags_list, cat))
    wb.close()

    if not rows:
        print("В xlsx нет подходящих строк (заполнены вопрос и ответ).")
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM knowledge_base")
            print(f"Записей в knowledge_base: {cur.fetchone()[0]}")
        return

    with conn.cursor() as cur:
        cur.execute("DELETE FROM knowledge_base")
        deleted = cur.rowcount
        print(f"Удалено предыдущих записей в knowledge_base: {deleted}")

        # Проверяем наличие колонки keywords
        cur.execute("""
            SELECT column_name FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = 'knowledge_base'
        """)
        cols = {r[0] for r in cur.fetchall()}
        has_keywords = "keywords" in cols
        has_embedding = "embedding" in cols

        def do_insert(row: tuple) -> None:
            t, c, s, tags, cat = row
            if has_embedding and has_keywords:
                cur.execute(
                    "INSERT INTO knowledge_base (title, content, short_answer, tags, category, keywords, embedding) VALUES (%s, %s, %s, %s, %s, %s, NULL)",
                    (t, c, s, tags, cat or None, tags),
                )
            elif has_embedding:
                cur.execute(
                    "INSERT INTO knowledge_base (title, content, short_answer, tags, category, embedding) VALUES (%s, %s, %s, %s, %s, NULL)",
                    (t, c, s, tags, cat or None),
                )
            elif has_keywords:
                cur.execute(
                    "INSERT INTO knowledge_base (title, content, short_answer, tags, category, keywords) VALUES (%s, %s, %s, %s, %s, %s)",
                    (t, c, s, tags, cat or None, tags),
                )
            else:
                cur.execute(
                    "INSERT INTO knowledge_base (title, content, short_answer, tags, category) VALUES (%s, %s, %s, %s, %s)",
                    (t, c, s, tags, cat or None),
                )

        for row in rows:
            do_insert(row)
        cur.execute("SELECT COUNT(*) FROM knowledge_base")
        total = cur.fetchone()[0]
    print(f"Загружено записей из kb_test.xlsx: {len(rows)}")
    print(f"Всего записей в knowledge_base: {total}")
    print("В колонке tags сохранены шаблонные вопросы (примеры запросов, приводящих к этой теме).")
    if has_embedding:
        print("Колонка embedding оставлена NULL — заполните через POST /kb/refresh-embeddings или скрипт.")


def exec_many(conn: PgConnection, statements: Iterable[str], title: str) -> None:
    print(f"\n== {title} ==")
    with conn.cursor() as cur:
        for idx, stmt in enumerate(statements, start=1):
            short = stmt.strip().splitlines()[0][:90]
            print(f"[{idx}] {short} ...")
            cur.execute(stmt)
    print("OK")


def ensure_database(
    host: str,
    port: int,
    user: str,
    password: str,
    db_name: str,
) -> None:
    host = str(host or "localhost").strip()
    user = str(user or "postgres").strip()
    password = str(password or "").strip()
    db_name = str(db_name or "postgres").strip()
    conn = psycopg.connect(
        host=host,
        port=port,
        user=user,
        password=password,
        dbname="postgres",
        options="-c client_encoding=UTF8",
    )
    conn.autocommit = True
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT 1 FROM pg_database WHERE datname = %s", (db_name,))
            if cur.fetchone():
                print(f"База {db_name!r} уже существует.")
                return

            cur.execute(
                sql.SQL(
                    "CREATE DATABASE {} WITH ENCODING 'UTF8' TEMPLATE template0"
                ).format(sql.Identifier(db_name))
            )
            print(f"Создана база {db_name!r}.")
    finally:
        conn.close()


def create_schema(
    host: str,
    port: int,
    user: str,
    password: str,
    db_name: str,
    drop_existing: bool,
    seed: bool,
) -> None:
    host = str(host or "localhost").strip()
    user = str(user or "postgres").strip()
    password = str(password or "").strip()
    db_name = str(db_name or "test").strip()
    conn = psycopg.connect(
        host=host,
        port=port,
        user=user,
        password=password,
        dbname=db_name,
        options="-c client_encoding=UTF8",
    )
    conn.autocommit = True

    try:
        if drop_existing:
            exec_many(
                conn,
                [
                    "DROP TABLE IF EXISTS feedback CASCADE;",
                    "DROP TABLE IF EXISTS email_log CASCADE;",
                    "DROP TABLE IF EXISTS knowledge_base CASCADE;",
                    "DROP TABLE IF EXISTS tickets CASCADE;",
                    "DROP TABLE IF EXISTS operators CASCADE;",
                ],
                "Удаление старых таблиц",
            )

        # Расширения: pg_trgm всегда, vector (pgvector) — опционально
        vector_available = False
        with conn.cursor() as cur:
            print("\n== Расширения ==")
            print("[1] CREATE EXTENSION IF NOT EXISTS pg_trgm ...")
            cur.execute("CREATE EXTENSION IF NOT EXISTS pg_trgm;")
            print("[2] CREATE EXTENSION IF NOT EXISTS vector ...")
            try:
                cur.execute("CREATE EXTENSION IF NOT EXISTS vector;")
                vector_available = True
                print("OK (pgvector)")
            except Exception as e:
                if "vector" in str(e).lower():
                    print("(пропущено: pgvector не установлен — колонка embedding не будет создана)")
                else:
                    raise

        # Если таблица tickets уже создана (например backend/db.py) без колонок tags/category/search_vector —
        # добавляем их до создания индексов (только если таблица уже есть).
        with conn.cursor() as cur:
            cur.execute("""
                SELECT 1 FROM information_schema.tables
                WHERE table_schema = 'public' AND table_name = 'tickets';
            """)
            if cur.fetchone():
                for col, typ in (
                    ("tags", "TEXT[]"),
                    ("category", "VARCHAR(100)"),
                    ("operator_id", "INTEGER REFERENCES operators(id) ON DELETE SET NULL"),
                    ("ai_processing_time", "INTEGER"),
                    ("ai_model", "VARCHAR(50)"),
                ):
                    try:
                        cur.execute(
                            f"ALTER TABLE tickets ADD COLUMN IF NOT EXISTS {col} {typ};"
                        )
                    except Exception:
                        pass
                try:
                    cur.execute("""
                        ALTER TABLE tickets ADD COLUMN IF NOT EXISTS search_vector tsvector
                        GENERATED ALWAYS AS (
                            setweight(to_tsvector('simple', coalesce(question,'')), 'A') ||
                            setweight(to_tsvector('simple', coalesce(answer,'')), 'B')
                        ) STORED;
                    """)
                except Exception:
                    pass

            # Если knowledge_base уже создана (например backend/db.py) без search_vector/keywords — добавляем.
            cur.execute("""
                SELECT 1 FROM information_schema.tables
                WHERE table_schema = 'public' AND table_name = 'knowledge_base';
            """)
            if cur.fetchone():
                for col, typ in (
                    ("keywords", "TEXT[]"),
                    ("success_rate", "FLOAT DEFAULT 1.0"),
                ):
                    try:
                        cur.execute(
                            f"ALTER TABLE knowledge_base ADD COLUMN IF NOT EXISTS {col} {typ};"
                        )
                    except Exception:
                        pass
                # search_vector с конфигом 'russian' для морфологии и стемминга
                try:
                    cur.execute("DROP INDEX IF EXISTS idx_kb_search;")
                    cur.execute("ALTER TABLE knowledge_base DROP COLUMN IF EXISTS search_vector;")
                    cur.execute("""
                        ALTER TABLE knowledge_base ADD COLUMN search_vector tsvector
                        GENERATED ALWAYS AS (
                            setweight(to_tsvector('russian', coalesce(title,'')), 'A') ||
                            setweight(to_tsvector('russian', coalesce(content,'')), 'B')
                        ) STORED;
                    """)
                    cur.execute("CREATE INDEX IF NOT EXISTS idx_kb_search ON knowledge_base USING GIN (search_vector);")
                except Exception as e:
                    print("(search_vector для knowledge_base:", e, ")")
                if vector_available:
                    try:
                        cur.execute(
                            "ALTER TABLE knowledge_base ADD COLUMN IF NOT EXISTS embedding vector(384);"
                        )
                    except Exception:
                        pass

        schema_statements = [
            """
            CREATE TABLE IF NOT EXISTS operators (
                id SERIAL PRIMARY KEY,
                email VARCHAR(255) UNIQUE NOT NULL,
                name VARCHAR(255) NOT NULL,
                is_active BOOLEAN DEFAULT TRUE,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            """,
            """
            CREATE TABLE IF NOT EXISTS tickets (
                id SERIAL PRIMARY KEY,
                client_email VARCHAR(255) NOT NULL,
                client_name VARCHAR(255),
                subject VARCHAR(500),
                question TEXT,
                answer TEXT,
                status VARCHAR(50) DEFAULT 'new',
                ai_confidence FLOAT,
                ai_suggested_answer TEXT,
                ai_category VARCHAR(100),
                ai_priority VARCHAR(50),
                ai_tone VARCHAR(50),
                ai_processing_time INTEGER,
                ai_model VARCHAR(50),
                tags TEXT[],
                category VARCHAR(100),
                operator_id INTEGER REFERENCES operators(id) ON DELETE SET NULL,
                needs_attention BOOLEAN DEFAULT FALSE,
                is_resolved BOOLEAN DEFAULT FALSE,
                message_id VARCHAR(255) UNIQUE,
                in_reply_to VARCHAR(255),
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                processed_at TIMESTAMP,
                resolved_at TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                search_vector tsvector GENERATED ALWAYS AS (
                    setweight(to_tsvector('simple', coalesce(question,'')), 'A') ||
                    setweight(to_tsvector('simple', coalesce(answer,'')), 'B')
                ) STORED
            );
            """,
            "CREATE INDEX IF NOT EXISTS idx_tickets_status ON tickets(status);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_client_email ON tickets(client_email);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_created_at ON tickets(created_at DESC);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_resolved_at ON tickets(resolved_at DESC);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_is_resolved ON tickets(is_resolved);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_needs_attention ON tickets(needs_attention);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_tags ON tickets USING GIN (tags);",
            "CREATE INDEX IF NOT EXISTS idx_tickets_search ON tickets USING GIN (search_vector);",
            """
            CREATE TABLE IF NOT EXISTS email_log (
                id SERIAL PRIMARY KEY,
                ticket_id INTEGER NOT NULL REFERENCES tickets(id) ON DELETE CASCADE,
                raw_from VARCHAR(255) NOT NULL,
                raw_to VARCHAR(255) NOT NULL,
                raw_subject VARCHAR(500),
                raw_body TEXT,
                raw_html TEXT,
                message_id VARCHAR(255),
                in_reply_to VARCHAR(255),
                direction VARCHAR(10) DEFAULT 'incoming',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            """,
            "CREATE INDEX IF NOT EXISTS idx_email_log_ticket ON email_log(ticket_id);",
            "CREATE INDEX IF NOT EXISTS idx_email_log_message_id ON email_log(message_id);",
            "CREATE INDEX IF NOT EXISTS idx_email_log_direction ON email_log(direction);",
            "CREATE INDEX IF NOT EXISTS idx_email_log_created ON email_log(created_at DESC);",
            (
                """
            CREATE TABLE IF NOT EXISTS knowledge_base (
                id SERIAL PRIMARY KEY,
                title VARCHAR(500) NOT NULL,
                content TEXT NOT NULL,
                short_answer TEXT,
                tags TEXT[],
                category VARCHAR(100),
                """
                + ("embedding vector(384),\n                " if vector_available else "")
                + """keywords TEXT[],
                usage_count INTEGER DEFAULT 0,
                success_rate FLOAT DEFAULT 1.0,
                is_active BOOLEAN DEFAULT TRUE,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                search_vector tsvector GENERATED ALWAYS AS (
                    setweight(to_tsvector('russian', coalesce(title,'')), 'A') ||
                    setweight(to_tsvector('russian', coalesce(content,'')), 'B')
                ) STORED
            );
            """,
            "CREATE INDEX IF NOT EXISTS idx_kb_tags ON knowledge_base USING GIN (tags);",
            "CREATE INDEX IF NOT EXISTS idx_kb_category ON knowledge_base(category);",
            "CREATE INDEX IF NOT EXISTS idx_kb_is_active ON knowledge_base(is_active);",
            "CREATE INDEX IF NOT EXISTS idx_kb_search ON knowledge_base USING GIN (search_vector);",
            "CREATE INDEX IF NOT EXISTS idx_kb_usage ON knowledge_base(usage_count DESC);",
            """
            CREATE TABLE IF NOT EXISTS feedback (
                id SERIAL PRIMARY KEY,
                ticket_id INTEGER NOT NULL REFERENCES tickets(id) ON DELETE CASCADE,
                operator_id INTEGER REFERENCES operators(id) ON DELETE SET NULL,
                kb_id INTEGER REFERENCES knowledge_base(id) ON DELETE SET NULL,
                is_helpful BOOLEAN,
                rating INTEGER CHECK (rating >= 1 AND rating <= 5),
                comment TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            );
            """,
            "CREATE INDEX IF NOT EXISTS idx_feedback_ticket ON feedback(ticket_id);",
            "CREATE INDEX IF NOT EXISTS idx_feedback_kb ON feedback(kb_id);",
            "CREATE INDEX IF NOT EXISTS idx_feedback_helpful ON feedback(is_helpful);",
        ]

        exec_many(conn, schema_statements, "Создание таблиц и индексов")

        # Индекс для векторного поиска (семантика)
        if vector_available:
            with conn.cursor() as cur:
                try:
                    cur.execute(
                        "CREATE INDEX IF NOT EXISTS idx_kb_embedding ON knowledge_base "
                        "USING ivfflat (embedding vector_cosine_ops) WITH (lists = 1);"
                    )
                    print("Индекс idx_kb_embedding создан (при заполнении БЗ пересоздайте с lists ≈ sqrt(rows)).")
                except Exception as e:
                    print("(idx_kb_embedding:", e, ")")

        # Добавить колонки, если таблица создана старым скриптом без них
        with conn.cursor() as cur:
            for col, typ in (
                ("ai_category", "VARCHAR(100)"),
                ("ai_priority", "VARCHAR(50)"),
                ("ai_tone", "VARCHAR(50)"),
                ("message_id", "VARCHAR(255)"),
                ("in_reply_to", "VARCHAR(255)"),
            ):
                try:
                    cur.execute(
                        f"ALTER TABLE tickets ADD COLUMN IF NOT EXISTS {col} {typ};"
                    )
                except Exception:
                    pass
            try:
                cur.execute(
                    "CREATE UNIQUE INDEX IF NOT EXISTS idx_tickets_message_id ON tickets (message_id) WHERE message_id IS NOT NULL;"
                )
            except Exception:
                pass

        if seed:
            seed_statements = [
                """
                INSERT INTO operators (email, name)
                VALUES ('operator@support.ru', 'Иван Петров'),
                       ('alex@support.ru', 'Алексей Смирнов')
                ON CONFLICT (email) DO NOTHING;
                """,
                """
                INSERT INTO tickets (
                    client_email, client_name, subject, question, answer, status, tags, category,
                    operator_id, needs_attention, is_resolved, created_at, resolved_at
                )
                SELECT
                    'ivanov@example.com',
                    'Иван Иванов',
                    'Насос X-100 не включается',
                    'Здравствуйте! Насос X-100 перестал запускаться. Мигает красная лампочка.',
                    'Проверьте предохранитель F2 на блоке питания.',
                    'resolved',
                    ARRAY['pump', 'electrical', 'X-100'],
                    'hardware',
                    (SELECT id FROM operators WHERE email = 'operator@support.ru'),
                    FALSE,
                    TRUE,
                    '2024-02-20 10:30:00',
                    '2024-02-20 11:15:00'
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM tickets
                    WHERE client_email = 'ivanov@example.com'
                      AND subject = 'Насос X-100 не включается'
                );
                """,
                """
                INSERT INTO tickets (
                    client_email, client_name, subject, question, status, tags, category,
                    operator_id, needs_attention, is_resolved
                )
                SELECT
                    'petrova@example.com',
                    'Мария Петрова',
                    'Ошибка E21 на котле',
                    'Добрый день. После включения через 2 минуты появляется ошибка E21.',
                    'in_progress',
                    ARRAY['boiler', 'E21'],
                    'hardware',
                    (SELECT id FROM operators WHERE email = 'alex@support.ru'),
                    TRUE,
                    FALSE
                WHERE NOT EXISTS (
                    SELECT 1
                    FROM tickets
                    WHERE client_email = 'petrova@example.com'
                      AND subject = 'Ошибка E21 на котле'
                );
                """,
                """
                INSERT INTO email_log (
                    ticket_id, raw_from, raw_to, raw_subject, raw_body, message_id, direction
                )
                SELECT
                    t.id,
                    'petrova@example.com',
                    'support@company.ru',
                    'Re: Ошибка E21 на котле',
                    'Пробовала перезапуск, ошибка осталась. Давление 0.8 bar.',
                    '<msg-e21-001@example.com>',
                    'incoming'
                FROM tickets t
                WHERE t.client_email = 'petrova@example.com'
                  AND t.subject = 'Ошибка E21 на котле'
                  AND NOT EXISTS (
                      SELECT 1
                      FROM email_log el
                      WHERE el.message_id = '<msg-e21-001@example.com>'
                  );
                """,
                """
                INSERT INTO feedback (
                    ticket_id, operator_id, kb_id, is_helpful, rating, comment
                )
                SELECT
                    t.id,
                    o.id,
                    kb.id,
                    TRUE,
                    5,
                    'Ответ помог, проблема решена быстро.'
                FROM tickets t
                JOIN operators o ON o.email = 'operator@support.ru'
                JOIN knowledge_base kb ON kb.title = 'Ошибка E21 на котле ThermoMax'
                WHERE t.client_email = 'ivanov@example.com'
                  AND t.subject = 'Насос X-100 не включается'
                  AND NOT EXISTS (
                      SELECT 1
                      FROM feedback f
                      WHERE f.ticket_id = t.id
                  );
                """,
            ]
            exec_many(conn, seed_statements, "Тестовые данные")

            # Заполнение knowledge_base из kb_test.xlsx (предыдущая вставка удаляется)
            _seed_kb_from_xlsx(conn)

        print("\nГотово. База и таблицы созданы.")
    finally:
        conn.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Создание БД support schema в PostgreSQL."
    )
    parser.add_argument("--host", default=os.getenv("PGHOST", "localhost"))
    parser.add_argument("--port", type=int, default=int(os.getenv("PGPORT", "5432")))
    parser.add_argument("--user", default=os.getenv("PGUSER", "postgres"))
    parser.add_argument(
        "--password",
        default=os.getenv("PGPASSWORD"),
        help="Пароль PostgreSQL (если не указан, будет запрошен интерактивно)",
    )
    parser.add_argument("--db", default=os.getenv("PGDATABASE", "test"), help="Имя целевой базы")
    parser.add_argument(
        "--create-db",
        action="store_true",
        help="Создать базу, если ее нет",
    )
    parser.add_argument(
        "--drop-existing",
        action="store_true",
        help="Удалить существующие таблицы перед созданием",
    )
    parser.add_argument(
        "--seed",
        action="store_true",
        help="Заполнить тестовыми данными",
    )

    args = parser.parse_args()
    if not args.password:
        args.password = getpass.getpass("Введите пароль PostgreSQL: ")
    # Пароль и остальные параметры — строго строка, без лишних символов (избегаем ошибки "missing = after P")
    args.password = str(args.password).strip() if args.password else ""
    args.host = str(args.host or "localhost").strip()
    args.user = str(args.user or "postgres").strip()
    args.db = str(args.db or os.getenv("PGDATABASE", "test")).strip()

    try:
        if args.create_db:
            ensure_database(args.host, args.port, args.user, args.password, args.db)

        create_schema(
            host=args.host,
            port=args.port,
            user=args.user,
            password=args.password,
            db_name=args.db,
            drop_existing=args.drop_existing,
            seed=args.seed,
        )
        return 0
    except Exception as exc:
        traceback.print_exc()
        print("\nОШИБКА:", exc)
        err_text = str(exc).lower()
        if "connection failed" in err_text or "connection refused" in err_text or "127.0.0.1" in str(exc):
            print("\n--- Подсказка: PostgreSQL недоступен по адресу localhost:5432.")
            print("  Инициализация из контейнера (из корня проекта):")
            print("    1) Пересоздать backend с монтированием проекта: docker compose up -d --force-recreate backend")
            print("    2) Выполнить: docker compose exec -e PGHOST=postgres backend python /project/init_database.py --create-db --seed")
            print("  Или с хоста: убедитесь, что postgres слушает 5432 (docker compose ps), затем снова эту команду.")
        elif "password authentication failed" in err_text or "authentication failed" in err_text:
            print("\n--- Подсказка: если пароль верный, попробуйте передать его в командной строке:")
            print('  python init_database.py --create-db --password "ВАШ_ПАРОЛЬ"')
            print("  Проверьте подключение вручную: psql -U postgres -h localhost")
            print("  Если не подходит — сбросьте пароль в PostgreSQL (pg_hba.conf или ALTER USER).")
        return 1


if __name__ == "__main__":
    sys.exit(main())
